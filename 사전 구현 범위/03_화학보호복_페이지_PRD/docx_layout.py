from pathlib import Path
from datetime import datetime
from zoneinfo import ZoneInfo
import hashlib
import json

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont


ROOT = Path('/home/giwon/Downloads/ppe-reference-poc')
OUT = Path(__file__).resolve().parent
OUT.mkdir(parents=True, exist_ok=True)
FILE = OUT / 'CHEMICAL_PPE_PAGE_PRD_20261009.docx'
FONT = 'Noto Sans CJK KR'
FONTFILE = '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'


def rf(run, size=None, bold=None, color='000000'):
    run.font.name = FONT
    props = run._element.get_or_add_rPr()
    fonts = props.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        props.insert(0, fonts)
    for field in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn('w:' + field), FONT)
    run.font.color.rgb = RGBColor.from_string(color)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(.66)
sec.bottom_margin = Inches(.66)
sec.left_margin = Inches(.72)
sec.right_margin = Inches(.72)
sec.footer_distance = Inches(.28)

for name, size in [('Normal', 11), ('Title', 24), ('Subtitle', 12),
                   ('Heading 1', 17), ('Heading 2', 12.5), ('Heading 3', 11)]:
    st = doc.styles[name]
    st.font.name = FONT
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.font.italic = False
    st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), FONT)
    pf = st.paragraph_format
    pf.line_spacing = Pt(16 if name == 'Normal' else size + 7)
    pf.space_after = Pt(6)
    if name.startswith('Heading'):
        st.font.bold = True
        pf.space_before = Pt(12 if name != 'Heading 1' else 0)
        pf.keep_with_next = True
doc.styles['Normal'].paragraph_format.widow_control = True

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
rf(footer.add_run('케미가드 PRD 1.0  |  '), 8.5, color='505050')
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)


def p(text='', bold=False, size=None):
    para = doc.add_paragraph()
    rf(para.add_run(text), size, bold)
    return para


def h(text, level=2):
    para = doc.add_paragraph(text, style=f'Heading {level}')
    for run in para.runs:
        rf(run)
    return para


def page(text):
    doc.add_page_break()
    h(text, 1)


def bullet(text):
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Inches(.12)
    para.paragraph_format.first_line_indent = Inches(-.12)
    rf(para.add_run('• ' + text))
    return para


def table(headers, rows, widths, size=10.2):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    props = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge = OxmlElement('w:' + side)
        edge.set(qn('w:val'), 'single')
        edge.set(qn('w:sz'), '4')
        edge.set(qn('w:color'), 'D9D9D9')
        borders.append(edge)
    props.append(borders)
    repeat = OxmlElement('w:tblHeader')
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for rownum, values in enumerate([headers] + rows):
        row = t.rows[0] if rownum == 0 else t.add_row()
        cant = OxmlElement('w:cantSplit')
        row._tr.get_or_add_trPr().append(cant)
        for i, text in enumerate(values):
            cell = row.cells[i]
            cell.width = Inches(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side in ('top', 'bottom'):
                el = OxmlElement('w:' + side)
                el.set(qn('w:w'), '90')
                el.set(qn('w:type'), 'dxa')
                margins.append(el)
            for side in ('left', 'right'):
                el = OxmlElement('w:' + side)
                el.set(qn('w:w'), '110')
                el.set(qn('w:type'), 'dxa')
                margins.append(el)
            tcpr.append(margins)
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), '243D4E' if rownum == 0 else ('F3F6F8' if rownum % 2 == 0 else 'FFFFFF'))
            tcpr.append(shade)
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = Pt(14.5)
            if i == 0 and len(headers) >= 3:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            rf(para.add_run(str(text)), size, rownum == 0, 'FFFFFF' if rownum == 0 else '000000')
    p('').paragraph_format.space_after = Pt(0)
    return t


def requirement(code, name, behavior, acceptance):
    h(f'{code} {name}', 2)
    p(behavior)
    para = p()
    rf(para.add_run('완료 기준  '), bold=True)
    rf(para.add_run(acceptance))


def link(text, target):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(4)
    rel = doc.part.relate_to(target, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    hyp = OxmlElement('w:hyperlink')
    hyp.set(qn('r:id'), rel)
    r = OxmlElement('w:r')
    pr = OxmlElement('w:rPr')
    fonts = OxmlElement('w:rFonts')
    for field in ('ascii', 'hAnsi', 'eastAsia'):
        fonts.set(qn('w:' + field), FONT)
    pr.append(fonts)
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '174B73')
    pr.append(color)
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), '20')
    pr.append(sz)
    r.append(pr)
    tt = OxmlElement('w:t')
    tt.text = text
    r.append(tt)
    hyp.append(r)
    para._p.append(hyp)


